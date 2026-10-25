"""Builds Dissertation_Data_Compilation.xlsx: every dataset used in the dissertation, in one workbook.
Raw data are entered as values (blue); every derived series (credit growth, growth gaps, share returns) is an
Excel formula (black) so that the workbook shows how each number is obtained.
Run: python 09_build_data_workbook.py   then recalculate (LibreOffice) before use.
"""
import glob
import os

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

HERE = os.path.dirname(os.path.abspath(__file__))
OUTF = os.path.join(HERE, "Dissertation_Data_Compilation.xlsx")
FONT = "Arial"
F_IN = Font(name=FONT, size=10, color="0000FF")          # raw input values
F_FX = Font(name=FONT, size=10, color="000000")          # formulas
F_HD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TX = Font(name=FONT, size=10)
F_TI = Font(name=FONT, size=12, bold=True)
F_NOTE = Font(name=FONT, size=9, italic=True, color="555555")
FILL_HD = PatternFill("solid", fgColor="1F3864")
FILL_SUB = PatternFill("solid", fgColor="D9E1F2")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()


def sheet(name, title, note):
    ws = wb.create_sheet(name)
    ws["A1"] = title; ws["A1"].font = F_TI
    ws["A2"] = note; ws["A2"].font = F_NOTE
    return ws


def header(ws, row, cols, widths=None):
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.font = F_HD; cell.fill = FILL_HD; cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center"); cell.border = BOX
    ws.row_dimensions[row].height = 42
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[L(j)].width = w


def put(ws, r, c, v, font=F_IN, fmt=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = font
    if fmt: cell.number_format = fmt
    return cell


# =================================================================== RBI data
raw = pd.read_excel(os.path.join(HERE, "RBI_Sectoral_Bank_Credit.xlsx"), header=None)
dates = pd.to_datetime(raw.iloc[6, 3:].astype(str).str.replace(r"\s+", " ", regex=True), format="%B %d, %Y")
ROWS = [("Non-food credit", 9, "III"), ("NBFCs (incl. HFCs)", 11, "3.9"), ("HFCs", 12, "3.9.1"), ("Housing", 13, "4.2"),
        ("Credit cards", 15, "4.5"), ("Education", 16, "4.6"), ("Vehicle loans", 17, "4.7"), ("Gold loans", 18, "4.8"), ("Other personal loans", 19, "4.9")]
EXR = [("Non-food credit (excl. merger)", 10), ("Housing (excl. merger)", 14), ("Other personal loans (excl. merger)", 20)]
lev = pd.DataFrame({k: pd.to_numeric(raw.iloc[r, 3:], errors="coerce").values for k, r, _ in ROWS}, index=dates)
exm = pd.DataFrame({k: pd.to_numeric(raw.iloc[r, 3:], errors="coerce").values for k, r in EXR}, index=dates)
N = len(dates); R0 = 5                                   # first data row
MERGER, EX_END = pd.Timestamp("2023-07-28"), exm.iloc[:, 1].last_valid_index()

# ---- RBI_Levels
ws = sheet("RBI_Levels", "Outstanding bank credit by sector, ₹ crore (reported figures)",
           "Source: RBI, Database on Indian Economy, Table 15 'Deployment of Gross Bank Credit by Major Sectors'. Each observation is the last reporting Friday of the month. Blue = data as published; black = formula.")
cols = ["Date"] + [f"{k} (RBI item {c})" for k, _, c in ROWS] + ["NBFCs excl. HFCs (formula)", "Non-food excl. NBFCs, cards and other PL (formula)"]
header(ws, 4, cols, [12] + [14] * 9 + [16, 20])
for i, d in enumerate(dates):
    r = R0 + i
    put(ws, r, 1, d.to_pydatetime(), F_IN, "dd-mmm-yyyy")
    for j, (k, _, _) in enumerate(ROWS, start=2):
        put(ws, r, j, float(lev.iloc[i][k]), F_IN, "#,##0.00")
    put(ws, r, 11, f"=C{r}-D{r}", F_FX, "#,##0.00")
    put(ws, r, 12, f"=B{r}-C{r}-F{r}-J{r}", F_FX, "#,##0.00")
ws.freeze_panes = "B5"
LV = {"Non-food credit": "B", "NBFCs (incl. HFCs)": "C", "HFCs": "D", "Housing": "E", "Credit cards": "F", "Education": "G", "Vehicle loans": "H",
      "Gold loans": "I", "Other personal loans": "J", "NBFCs ex-HFC": "K", "Core": "L"}

# ---- RBI_Levels_ExMerger
ws = sheet("RBI_Levels_ExMerger", "Bank credit excluding the impact of the HDFC merger, ₹ crore (July 2023 to June 2025)",
           "Source: RBI DBIE Table 15, figures shown in parentheses (note 5: 'Data since July 28, 2023 include the impact of the merger of a non-bank with a bank. Figures in parentheses exclude the impact of the merger.'). Blank where not published.")
header(ws, 4, ["Date", "Non-food credit (excl. merger)", "Housing (excl. merger)", "Other personal loans (excl. merger)",
               "Non-food excl. NBFCs, cards and other PL (excl. merger, formula)"], [12, 18, 18, 20, 24])
for i, d in enumerate(dates):
    r = R0 + i
    put(ws, r, 1, d.to_pydatetime(), F_IN, "dd-mmm-yyyy")
    for j, (k, _) in enumerate(EXR, start=2):
        v = exm.iloc[i][k]
        if not np.isnan(v): put(ws, r, j, float(v), F_IN, "#,##0.00")
    put(ws, r, 5, f'=IF(B{r}="","",B{r}-RBI_Levels!C{r}-RBI_Levels!F{r}-D{r})', F_FX, "#,##0.00")
ws.freeze_panes = "B5"
EX = {"Non-food credit": "B", "Housing": "C", "Other personal loans": "D", "Core": "E"}

# ---- RBI_Growth: monthly log growth x 100 (formulas)
ws = sheet("RBI_Growth", "Monthly growth of bank credit (100 × change in natural log) and the growth gaps used in the tests",
           "Formulas. Merger month (28-Jul-2023) left blank for merger-affected series. From Aug-2023 to Jun-2025, growth of non-food credit, housing, other personal loans and the alternative benchmark uses the figures excluding the merger.")
GSER = ["Credit cards", "Other personal loans", "Housing", "Vehicle loans", "Education", "Gold loans", "NBFCs ex-HFC", "HFCs", "NBFCs (incl. HFCs)", "Non-food credit", "Core"]
GNAME = {"Core": "Non-food excl. NBFCs, cards, other PL"}
MERGED = {"Housing", "Other personal loans", "NBFCs (incl. HFCs)", "HFCs", "NBFCs ex-HFC", "Non-food credit", "Core"}
gcols = ["Date"] + [GNAME.get(s, s) for s in GSER] + ["Gap1: (cards + other PL)/2 minus (housing + vehicle + education)/3",
                                                    "Gap2: NBFCs ex-HFC minus non-food credit", "Gap2 (alt.): NBFCs ex-HFC minus non-food excl. NBFCs, cards, PL",
                                                    "Test window (Apr-21 to Jan-25)", "Post (1 from 29-Dec-2023)", "Period"]
header(ws, 4, gcols, [12] + [12] * len(GSER) + [20, 16, 20, 12, 10, 16])
gcol = {s: L(2 + k) for k, s in enumerate(GSER)}
for i, d in enumerate(dates):
    r = R0 + i
    put(ws, r, 1, d.to_pydatetime(), F_IN, "dd-mmm-yyyy")
    if i == 0:
        continue
    for k, s in enumerate(GSER):
        c = 2 + k
        if d == MERGER and s in MERGED:
            put(ws, r, c, "", F_FX); ws.cell(row=r, column=c).comment = Comment("Merger month: excluded (HDFC Ltd merged into HDFC Bank).", "Note") if s == "Housing" else None
            continue
        use_ex = s in EX and MERGER < d <= EX_END
        src, col = ("RBI_Levels_ExMerger", EX[s]) if use_ex else ("RBI_Levels", LV[s])
        put(ws, r, c, f"=100*LN({src}!{col}{r}/{src}!{col}{r - 1})", F_FX, "0.000")
    a, b = gcol, None
    put(ws, r, 2 + len(GSER), f'=AVERAGE({a["Credit cards"]}{r},{a["Other personal loans"]}{r})-AVERAGE({a["Housing"]}{r},{a["Vehicle loans"]}{r},{a["Education"]}{r})', F_FX, "0.000")
    put(ws, r, 3 + len(GSER), f'=IF(OR({a["NBFCs ex-HFC"]}{r}="",{a["Non-food credit"]}{r}=""),"",{a["NBFCs ex-HFC"]}{r}-{a["Non-food credit"]}{r})', F_FX, "0.000")
    put(ws, r, 4 + len(GSER), f'=IF(OR({a["NBFCs ex-HFC"]}{r}="",{a["Core"]}{r}=""),"",{a["NBFCs ex-HFC"]}{r}-{a["Core"]}{r})', F_FX, "0.000")
    put(ws, r, 5 + len(GSER), f'=IF(AND(A{r}>=DATE(2021,4,1),A{r}<=DATE(2025,1,31)),1,0)', F_FX, "0")
    put(ws, r, 6 + len(GSER), f'=IF(A{r}>=DATE(2023,12,1),1,0)', F_FX, "0")
    put(ws, r, 7 + len(GSER), f'=IF(A{r}<DATE(2021,4,1),"Before sample",IF(A{r}<=DATE(2023,11,17),"Pre-policy",IF(A{r}<=DATE(2025,3,31),"Post-policy","After rollback")))', F_FX)
ws.freeze_panes = "B5"

# =================================================================== share prices
P = {}
for f in sorted(glob.glob(os.path.join(HERE, "event_study", "raw", "*.csv"))):
    d = pd.read_csv(f, encoding="utf-8-sig")
    d["Date"] = pd.to_datetime(d.Date, format="%d-%m-%Y")
    d["Price"] = pd.to_numeric(d.Price.astype(str).str.replace(",", ""))
    P[os.path.basename(f)[:-4]] = d.set_index("Date").Price.sort_index()
P = pd.DataFrame(P)
ORDER = ["NIFTY50", "NIFTY500", "AAVAS", "CANFINHOME", "HOMEFIRST", "LICHSGFIN", "PNBHOUSING", "REPCOHOME",
         "ARMANFIN", "BAJFINANCE", "CHOLAFIN", "CREDITACC", "IIFL", "MMFIN", "MUTHOOTFIN", "POONAWALLA", "SBICARD", "SHRIRAMFIN"]
P = P[ORDER]
ws = sheet("Share_Prices", "Daily closing prices (₹; index points for Nifty), 1 September 2022 to 28 March 2025",
           "Source: Investing.com historical data (NSE listings), adjusted for splits and bonus issues. 26-Feb-2025 was a market holiday (Mahashivratri).")
header(ws, 4, ["Date"] + ORDER, [12] + [12] * len(ORDER))
for i, (d, row) in enumerate(P.iterrows()):
    r = R0 + i
    put(ws, r, 1, d.to_pydatetime(), F_IN, "dd-mmm-yyyy")
    for j, s in enumerate(ORDER, start=2):
        put(ws, r, j, float(row[s]), F_IN, "#,##0.00")
ws.freeze_panes = "B5"
ws = sheet("Share_Returns", "Daily simple returns (formula: today's close / previous close − 1)",
           "Formulas referencing Share_Prices. Used in the event study.")
header(ws, 4, ["Date"] + ORDER, [12] + [12] * len(ORDER))
for i, d in enumerate(P.index):
    r = R0 + i
    put(ws, r, 1, f"=Share_Prices!A{r}", F_FX, "dd-mmm-yyyy")
    if i == 0: continue
    for j in range(2, len(ORDER) + 2):
        c = L(j)
        put(ws, r, j, f"=Share_Prices!{c}{r}/Share_Prices!{c}{r - 1}-1", F_FX, "0.00%")
ws.freeze_panes = "B5"

# =================================================================== firm data
ver = pd.read_csv(os.path.join(HERE, "data_corrections", "bank_exposure_verified.csv"))
SYM = {"AAVAS": "AAVAS", "Arman": "ARMANFIN", "Bajaj": "BAJFINANCE", "Can Fin": "CANFINHOME", "Cholamandalam": "CHOLAFIN", "CreditAccess": "CREDITACC",
       "Home First": "HOMEFIRST", "IIFL": "IIFL", "LIC Housing": "LICHSGFIN", "Mahindra": "MMFIN", "Muthoot": "MUTHOOTFIN", "PNB Housing": "PNBHOUSING",
       "Poonawalla": "POONAWALLA", "Repco": "REPCOHOME", "SBI Cards": "SBICARD", "Shriram": "SHRIRAMFIN"}
HFCK = ["AAVAS", "Can Fin", "Home First", "LIC Housing", "PNB Housing", "Repco"]
sym = lambda c: next(v for k, v in SYM.items() if k in c)
ws = sheet("Firm_Sample", "Sample of 16 listed NBFCs and HFCs", "Bank share at 30-Sep-2023 from each company's Q2 FY2024 investor presentation (Poonawalla does not disclose its borrowing mix).")
header(ws, 4, ["Company", "NSE symbol", "Type", "Group used in tests", "Consumer lender", "Microfinance lender", "Bank share of borrowings, 30-Sep-2023 (%)", "Definition in filing"],
       [40, 13, 14, 16, 11, 12, 16, 60])
for i, rw in enumerate(ver.sort_values("Company").itertuples()):
    r = R0 + i; c = rw.Company; s = sym(c)
    typ = "HFC" if any(h in c for h in HFCK) else {"CREDITACC": "NBFC (MFI)", "ARMANFIN": "NBFC (MFI)", "MUTHOOTFIN": "NBFC (gold)", "SBICARD": "NBFC (cards)"}.get(s, "NBFC")
    vals = [c, s, typ, "HFC (exempt)" if typ == "HFC" else "Non-HFC NBFC", "Yes" if s in ("BAJFINANCE", "SBICARD", "POONAWALLA") else "No",
            "Yes" if s in ("CREDITACC", "ARMANFIN") else "No", None if s == "POONAWALLA" else float(rw.Verified_Bank_Share_Sep2023_Pct),
            "Not disclosed" if s == "POONAWALLA" else rw.Definition_in_filing]
    for j, v in enumerate(vals, start=1):
        put(ws, r, j, v, F_IN if j == 7 else F_TX, "0.0" if j == 7 else None)
ws.freeze_panes = "B5"

for nm, fn, col, extra, title in [("Firm_Cost_of_Funds", "cof_quarterly.csv", "COF_Pct", "Note", "Cost of funds by company and quarter (%)"),
                                  ("Firm_Bank_Share", "bank_share_quarterly.csv", "Bank_Share_Pct", "Definition", "Bank share of borrowings by company and quarter (%)")]:
    d = pd.read_csv(os.path.join(HERE, "hand_collected", fn))
    d["t"] = (d.FY.str[2:].astype(int) - 21) * 4 + d.Q.str[1].astype(int) - 1
    d = d.sort_values(["Company", "t"])
    ws = sheet(nm, title, "Source: company quarterly investor presentations. 'Source' gives the presentation (FY-quarter) and page. Period: Pre = up to Q2 FY2024; Post = Q3 FY2024 to Q4 FY2025 (before the rollback); Later = FY2026.")
    header(ws, 4, ["Company", "NSE symbol", "Group", "Financial year", "Quarter", "Value (%)", "Period", "Source (presentation, page)", "Note / definition"],
           [40, 13, 14, 10, 9, 10, 9, 26, 70])
    for i, rw in enumerate(d.itertuples()):
        r = R0 + i
        grp = "HFC (exempt)" if any(h in rw.Company for h in HFCK) else "Non-HFC NBFC"
        vals = [rw.Company, sym(rw.Company), grp, rw.FY, rw.Q, float(getattr(rw, col)), None, rw.Source_deck, getattr(rw, extra) if isinstance(getattr(rw, extra), str) else ""]
        for j, v in enumerate(vals, start=1):
            put(ws, r, j, v, F_IN if j == 6 else F_TX, "0.00" if j == 6 else None)
        # period from FY / quarter (formula)
        put(ws, r, 7, f'=IF(VALUE(MID(D{r},3,2))*10+VALUE(MID(E{r},2,1))<=242,"Pre",IF(VALUE(MID(D{r},3,2))*10+VALUE(MID(E{r},2,1))<=254,"Post","Later"))', F_FX)
    ws.freeze_panes = "B5"

seg = pd.read_csv(os.path.join(HERE, "hand_collected", "segments_bajaj.csv"))
ws = sheet("Bajaj_Segments", "Bajaj Finance: assets under management by business segment (₹ crore)",
           "Source: Bajaj Finance quarterly investor presentations. Affected consumer segments: urban and rural sales finance and personal-loan (B2C) segments.")
header(ws, 4, ["Financial year", "Quarter", "Total AUM (₹ cr)", "Affected consumer segments (₹ cr)", "Share affected (formula)", "Source (presentation, page)", "Segment detail as reported"],
       [10, 9, 14, 16, 12, 18, 120])
for i, rw in enumerate(seg.itertuples()):
    r = R0 + i
    for j, v in enumerate([rw.FY, rw.Q, float(rw.Total_AUM_Cr), float(rw.Affected_Consumer_AUM_Cr)], start=1):
        put(ws, r, j, v, F_IN if j > 2 else F_TX, "#,##0" if j > 2 else None)
    put(ws, r, 5, f"=D{r}/C{r}", F_FX, "0.0%"); put(ws, r, 6, rw.Source, F_TX); put(ws, r, 7, rw.Segments, F_TX)
ws.freeze_panes = "A5"

ws = sheet("Company_Disclosures", "Company statements on the effect of the November 2023 measures", "Verbatim quotes from investor presentations for the quarter ended 31 December 2023.")
header(ws, 4, ["Company", "Source", "Statement (verbatim)", "Implication"], [26, 40, 90, 28])
for i, v in enumerate([
    ("Bajaj Finance Ltd.", "Q3 FY24 investor presentation", "RBI increased risk weights on consumer credit exposure from 100% to 125% which had an impact of 290 bps on the Company's CRAR. Adjusted for this change CRAR would have been 26.77%.", "Capital ratio cut by 2.9 pp"),
    ("Bajaj Finance Ltd.", "Q3 FY24 investor presentation", "Given the increase in risk weights and higher incremental cost of funds, the Company has increased rates across all portfolios by 20-30 bps.", "Lending rates raised 20-30 bps"),
    ("SBI Cards and Payment Services Ltd.", "Q3 FY24 investor presentation, capital adequacy slide", "CRAR impacted by ~400 bps due to increase in risk weight by RBI", "Capital ratio cut by about 4 pp")]):
    for j, x in enumerate(v, start=1):
        c = put(ws, R0 + i, j, x, F_TX); c.alignment = WRAP

ws = sheet("Event_Dates", "Policy events and the dates used in the analysis", "Release times are from the earliest dated news reports; NSE trading hours are 9:15 am to 3:30 pm IST.")
header(ws, 4, ["Event", "Date", "Time / note", "Source", "Used as"], [44, 13, 44, 60, 30])
for i, v in enumerate([
    ("RBI circular RBI/2023-24/85 (risk weights, exposure limits)", "2023-11-16", "First news report 5:26 pm IST, after market close", "RBI circular; Business Today, 16 Nov 2023", "Policy date"),
    ("First trading day after the circular", "2023-11-17", "Day 0 of event 1", "NSE trading calendar", "Event study day 0"),
    ("RBI embargo on Bajaj Finance eCOM and Insta EMI Card", "2023-11-15", "Lifted 2 May 2024", "Company announcements", "Confounding event (Bajaj excluded in checks)"),
    ("Deadline for board-approved limits on unsecured consumer credit", "2024-02-29", "", "RBI circular RBI/2023-24/85", "Context"),
    ("RBI embargo on IIFL Finance gold loans", "2024-03-04", "Lifted 19 Sep 2024", "Company announcements", "Context"),
    ("RBI circulars RBI/2024-25/119 and /120 (rollback of bank-to-NBFC risk weight)", "2025-02-25", "First news report 9:31 pm IST, after market close", "RBI circulars; Business Standard, 25 Feb 2025", "Rollback announcement"),
    ("Market holiday (Mahashivratri)", "2025-02-26", "NSE closed", "NSE holiday circular 170/2024", ""),
    ("First trading day after the rollback announcement", "2025-02-27", "Day 0 of event 2", "NSE trading calendar", "Event study day 0"),
    ("Rollback effective", "2025-04-01", "", "RBI circular RBI/2024-25/120", "Start of after-rollback period"),
    ("Bank credit tests: pre-policy period", "2021-04-01", "to 17-Nov-2023 (32 months)", "RBI DBIE Table 15", "Before"),
    ("Bank credit tests: post-policy period", "2023-12-29", "to 31-Jan-2025 (14 months)", "RBI DBIE Table 15", "After"),
    ("Event study estimation window, event 1", "2022-11-16", "to 04-Oct-2023 (220 trading days)", "Share_Prices", "Market model"),
    ("Event study estimation window, event 2", "2024-02-27", "to 15-Jan-2025 (220 trading days)", "Share_Prices", "Market model")]):
    for j, x in enumerate(v, start=1):
        c = put(ws, R0 + i, j, pd.Timestamp(x).to_pydatetime() if j == 2 else x, F_TX, "dd-mmm-yyyy" if j == 2 else None); c.alignment = WRAP

# =================================================================== README (first sheet)
ws = wb["Sheet"]; ws.title = "README"
ws["A1"] = "Dissertation data compilation"; ws["A1"].font = Font(name=FONT, size=14, bold=True)
ws["A2"] = "Capital Requirements and Credit Supply: The Effect of the RBI's November 2023 Risk-Weight Measures on Bank Consumer Credit and Bank Funding of NBFCs in India, 2021–2026"
ws["A2"].font = F_NOTE
header(ws, 4, ["Sheet", "Contents", "Frequency / coverage", "Source"], [24, 80, 30, 50])
info = [
    ("RBI_Levels", "Outstanding bank credit by sector (₹ crore): non-food credit, NBFCs, HFCs, housing, credit cards, education, vehicle loans, gold loans, other personal loans; NBFCs excl. HFCs and the alternative H2 benchmark by formula", "Monthly, Jan 2019 to Jul 2026 (91 observations)", "RBI DBIE Table 15"),
    ("RBI_Levels_ExMerger", "The same series excluding the impact of the HDFC merger, as published by the RBI", "Monthly, Jul 2023 to Jun 2025", "RBI DBIE Table 15 (figures in parentheses)"),
    ("RBI_Growth", "Monthly growth (100 × log change) of each series and the growth gaps used for H1 and H2, with window, post-policy and period flags (all formulas)", "Monthly", "Calculated"),
    ("Share_Prices", "Daily closing prices of the 16 sample companies and the Nifty 50 and Nifty 500", "Daily, 1 Sep 2022 to 28 Mar 2025 (641 trading days)", "Investing.com"),
    ("Share_Returns", "Daily simple returns (formulas)", "Daily", "Calculated"),
    ("Firm_Sample", "Company list, NSE symbol, type, groups used in the tests and pre-policy bank dependence", "30 Sep 2023", "Q2 FY2024 investor presentations"),
    ("Firm_Cost_of_Funds", "Cost of funds by company and quarter, with source presentation and page, and the pre/post period (formula)", "Quarterly, Q1 FY2021 to Q4 FY2026", "Investor presentations"),
    ("Firm_Bank_Share", "Bank share of borrowings by company and quarter, with source presentation and page, definition and period", "Quarterly, Q1 FY2021 to Q4 FY2026", "Investor presentations"),
    ("Bajaj_Segments", "Bajaj Finance assets under management by segment, and the share in segments affected by the circular (formula)", "Quarterly, Q1 FY2021 to Q4 FY2026", "Bajaj Finance investor presentations"),
    ("Company_Disclosures", "Verbatim company statements on the effect of the measures", "Q3 FY2024", "Investor presentations"),
    ("Event_Dates", "Policy and firm events, release times and the dates and windows used in the analysis", "", "RBI circulars, news reports, NSE"),
]
for i, v in enumerate(info):
    for j, x in enumerate(v, start=1):
        c = put(ws, R0 + i, j, x, F_TX); c.alignment = WRAP
r = R0 + len(info) + 1
for t in ["Colour code: blue = data as published or as compiled from the source; black = formula.",
          "Units: bank credit in ₹ crore (1 lakh crore = 100,000 crore); growth in percentage points per month; returns in per cent.",
          "Merger adjustment: HDFC Ltd merged into HDFC Bank on 1 July 2023. The merger month (28 Jul 2023) is excluded from growth for affected series, and growth from Aug 2023 to Jun 2025 uses the RBI figures that exclude the merger."]:
    ws.cell(row=r, column=1, value=t).font = F_TX; r += 1
ws.freeze_panes = "A5"
for w in wb.worksheets:
    w.sheet_view.zoomScale = 90
wb.save(OUTF)
print("saved", OUTF)
