"""Recalculate every formula in an .xlsx with LibreOffice (UNO) and save it in place, then report any error cells.
Run: python3 tools_xlsx_recalc.py Workbook.xlsx"""
import os, subprocess, sys, time, uno
from com.sun.star.beans import PropertyValue
from openpyxl import load_workbook


def prop(n, v):
    p = PropertyValue(); p.Name, p.Value = n, v; return p


path = os.path.abspath(sys.argv[1])
proc = subprocess.Popen(["soffice", "--headless", "--invisible", "--norestore", "--nologo",
                         "--accept=socket,host=localhost,port=2003;urp;"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
ctx = uno.getComponentContext()
resolver = ctx.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", ctx)
for _ in range(60):
    try:
        rctx = resolver.resolve("uno:socket,host=localhost,port=2003;urp;StarOffice.ComponentContext"); break
    except Exception:
        time.sleep(1)
desktop = rctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", rctx)
doc = desktop.loadComponentFromURL(uno.systemPathToFileUrl(path), "_blank", 0, (prop("Hidden", True),))
doc.calculateAll()
doc.storeToURL(uno.systemPathToFileUrl(path), (prop("FilterName", "Calc MS Excel 2007 XML"), prop("Overwrite", True)))
doc.close(True)
try:
    desktop.terminate()
except Exception:
    pass
proc.wait(timeout=30)

f, v = load_workbook(path), load_workbook(path, data_only=True)
n_f, errs = 0, []
for ws in f.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                n_f += 1
                val = v[ws.title][c.coordinate].value
                if isinstance(val, str) and val.startswith("#"):
                    errs.append(f"{ws.title}!{c.coordinate} {val}")
print({"total_formulas": n_f, "total_errors": len(errs), "errors": errs[:50]})
